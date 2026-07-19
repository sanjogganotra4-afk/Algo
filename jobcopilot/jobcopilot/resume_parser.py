"""Parse a .docx CV into structured JSON and cache it.

Usage:
    python -m jobcopilot.resume_parser /path/to/cv.docx
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from . import config, llm


def extract_docx_text(path: str | Path) -> str:
    """Extract raw text from a .docx CV (paragraphs + tables)."""
    from docx import Document  # python-docx

    doc = Document(str(path))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_and_cache(cv_path: str | Path) -> dict[str, Any]:
    """Parse the CV, structure it (LLM or heuristic), cache to resume_structured.json."""
    cv_path = Path(cv_path).expanduser()
    if not cv_path.exists():
        raise FileNotFoundError(f"CV not found: {cv_path}")
    raw = extract_docx_text(cv_path)
    if not raw.strip():
        raise ValueError("No text extracted from the .docx — is it empty or scanned?")
    structured = llm.structure_resume(raw)
    structured["_cv_path"] = str(cv_path)
    config.RESUME_STRUCTURED_PATH.write_text(
        json.dumps(structured, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    return structured


def load_cached() -> dict[str, Any]:
    if config.RESUME_STRUCTURED_PATH.exists():
        try:
            return json.loads(config.RESUME_STRUCTURED_PATH.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            pass
    return {"skills": [], "summary": "", "_source": "none"}


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python -m jobcopilot.resume_parser /path/to/cv.docx")
        sys.exit(1)
    result = parse_and_cache(sys.argv[1])
    src = result.get("_source")
    print(f"Parsed CV ({src}). Skills found: {len(result.get('skills', []))}")
    print(f"Cached to {config.RESUME_STRUCTURED_PATH}")


if __name__ == "__main__":
    main()
